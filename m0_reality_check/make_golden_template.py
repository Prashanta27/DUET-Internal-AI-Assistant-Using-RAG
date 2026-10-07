#!/usr/bin/env python3
"""Builds golden_questions.xlsx -- the M0 evaluation set skeleton."""

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT = "golden_questions.xlsx"
FIRST_DATA_ROW = 5          # rows 2-4 are worked examples
LAST_DATA_ROW = 104         # room for 100 questions

BASE_FONT = "Arial"
BN_FONT = "Nirmala UI"      # renders Bangla + Latin on Windows; Excel falls back elsewhere

HEAD_FILL = PatternFill("solid", fgColor="1F3864")
HEAD_FONT = Font(name=BASE_FONT, bold=True, color="FFFFFF", size=10)
EX_FILL = PatternFill("solid", fgColor="E2EFDA")
FILL_ME = PatternFill("solid", fgColor="FFFF00")
TITLE_FONT = Font(name=BASE_FONT, bold=True, size=13, color="1F3864")
BOLD = Font(name=BASE_FONT, bold=True, size=10)
BODY = Font(name=BASE_FONT, size=10)
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

COLUMNS = [
    ("id",                "Row number. Pre-filled -- do not change.", 6),
    ("language",          "bn = Bangla script | en = English | bnglish = romanised Bangla", 11),
    ("query_type",        "Pick from the dropdown.", 14),
    ("question",          "Type the question EXACTLY as a teacher would ask it.", 42),
    ("expected_answer",   "The short correct answer. A date, a name, a number, a rule.", 38),
    ("expected_document", "Exact filename the answer comes from, e.g. exam_notice_2026.pdf", 26),
    ("expected_page",     "Page number in that document.", 10),
    ("expected_table_id", "Only for table questions. Leave blank otherwise.", 12),
    ("notes",             "Anything tricky about this question.", 26),
]

QUERY_TYPES = [
    "factual", "date", "policy", "name", "number",
    "table", "multi_hop", "comparison", "unanswerable",
]

EXAMPLES = [
    ["EX1", "bn", "date",
     "২০২৬ সালের CSE পরীক্ষার রুটিন কবে প্রকাশিত হবে?",
     "১২ মার্চ ২০২৬",
     "academic_calendar_2026.pdf", 4, "",
     "Bangla numerals in the source -- check the answer is not converted to 12"],
    ["EX2", "en", "table",
     "Show the CSE first-year courses with their credits as a table.",
     "CSE1101 Programming 3; CSE1102 Data Structure 3",
     "syllabus_cse_2026.pdf", 12, "T2",
     "Whole table must come back intact, not split across chunks"],
    ["EX3", "bnglish", "factual",
     "thesis submission er last date kobe?",
     "30 June 2026",
     "thesis_guideline_2026.pdf", 7, "",
     "Romanised Bangla -- this is the query-normalisation test case"],
]

INSTRUCTIONS = [
    ("What this file is", None),
    ("The golden evaluation set for the DUET teaching assistant. Every retrieval and", None),
    ("model decision in this project gets judged against these questions. Without it", None),
    ("you are tuning blind and cannot report a single honest number.", None),
    ("", None),
    ("How to fill it in", None),
    ("1. Open the 'golden_questions' sheet. Rows 2-4 are worked examples -- read them,", None),
    ("   then DELETE them before you report results.", None),
    ("2. Start at row 5. Yellow cells are the ones you fill in.", None),
    ("3. Write questions a real DUET teacher would actually type. Not questions you", None),
    ("   already know the system can answer. Biased questions produce a useless score.", None),
    ("4. Open the source document, find the answer, and record the exact filename and", None),
    ("   page. If you cannot find it in a document, the question does not belong here.", None),
    ("5. Include roughly 15% 'unanswerable' questions -- things genuinely NOT in any", None),
    ("   document. Leave expected_answer blank for those. The system must refuse them.", None),
    ("   A model that never refuses will look perfect on the other 85% and be useless.", None),
    ("", None),
    ("Target for M0", None),
    ("30 questions: 10 Bangla, 10 English, 10 Banglish. That is enough to steer every", None),
    ("early decision. Expand to 150 at M5, not before -- writing 300 up front will", None),
    ("burn a week you need for the pipeline.", None),
    ("", None),
    ("Rule", None),
    ("Write the question and the expected answer BEFORE you run the system on it.", None),
    ("Writing them afterwards, from what the system happened to output, is not", None),
    ("evaluation -- it is decoration, and your adviser will see through it.", None),
]


def main() -> None:
    wb = Workbook()

    # ---------------- Instructions sheet ----------------
    ins = wb.active
    ins.title = "Instructions"
    ins["A1"] = "DUET Teaching Assistant -- Golden Evaluation Set"
    ins["A1"].font = TITLE_FONT

    r = 3
    for text, _ in INSTRUCTIONS:
        ins.cell(row=r, column=1, value=text)
        heading = text and not text.startswith((" ", "1", "2", "3", "4", "5")) and len(text) < 40
        ins.cell(row=r, column=1).font = BOLD if heading else BODY
        r += 1

    # Live coverage counters
    c = r + 1
    ins.cell(row=c, column=1, value="Coverage (updates automatically)").font = BOLD
    c += 1
    ins.cell(row=c, column=1, value="Bucket").font = HEAD_FONT
    ins.cell(row=c, column=2, value="Written").font = HEAD_FONT
    ins.cell(row=c, column=3, value="M0 target").font = HEAD_FONT
    ins.cell(row=c, column=4, value="Still to write").font = HEAD_FONT
    for col in range(1, 5):
        ins.cell(row=c, column=col).fill = HEAD_FILL
    header_row = c

    rng = f"golden_questions!$B${FIRST_DATA_ROW}:$B${LAST_DATA_ROW}"
    buckets = [("Bangla (bn)", "bn", 10), ("English (en)", "en", 10),
               ("Banglish (bnglish)", "bnglish", 10)]
    c += 1
    for label, code, target in buckets:
        ins.cell(row=c, column=1, value=label).font = BODY
        ins.cell(row=c, column=2, value=f'=COUNTIF({rng},"{code}")').font = BODY
        ins.cell(row=c, column=3, value=target).font = BODY
        ins.cell(row=c, column=4, value=f"=MAX(0,C{c}-B{c})").font = BODY
        c += 1

    qrng = f"golden_questions!$C${FIRST_DATA_ROW}:$C${LAST_DATA_ROW}"
    ins.cell(row=c, column=1, value="of which 'unanswerable'").font = BODY
    ins.cell(row=c, column=2, value=f'=COUNTIF({qrng},"unanswerable")').font = BODY
    ins.cell(row=c, column=3, value=5).font = BODY
    ins.cell(row=c, column=4, value=f"=MAX(0,C{c}-B{c})").font = BODY
    c += 1
    ins.cell(row=c, column=1, value="TOTAL written").font = BOLD
    ins.cell(row=c, column=2,
             value=f'=COUNTA(golden_questions!$D${FIRST_DATA_ROW}:$D${LAST_DATA_ROW})').font = BOLD
    ins.cell(row=c, column=3, value=30).font = BOLD
    ins.cell(row=c, column=4, value=f"=MAX(0,C{c}-B{c})").font = BOLD
    for rr in range(header_row, c + 1):
        for cc in range(1, 5):
            ins.cell(row=rr, column=cc).border = BOX

    ins.column_dimensions["A"].width = 82
    for col in "BCD":
        ins.column_dimensions[col].width = 14

    # ---------------- Data sheet ----------------
    ws = wb.create_sheet("golden_questions")
    for i, (name, helptext, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=i, value=name)
        cell.font = HEAD_FONT
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.comment = None
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[1].height = 22

    bn_cols = {4, 5}  # question, expected_answer may contain Bangla

    for j, row in enumerate(EXAMPLES):
        rr = 2 + j
        for i, val in enumerate(row, start=1):
            cell = ws.cell(row=rr, column=i, value=val)
            cell.font = Font(name=BN_FONT if i in bn_cols else BASE_FONT,
                             size=10, italic=True)
            cell.fill = EX_FILL
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = BOX
        ws.row_dimensions[rr].height = 30

    for rr in range(FIRST_DATA_ROW, LAST_DATA_ROW + 1):
        idx = rr - FIRST_DATA_ROW + 1
        ws.cell(row=rr, column=1, value=idx).font = BODY
        ws.cell(row=rr, column=1).alignment = Alignment(horizontal="center")
        for i in range(2, len(COLUMNS) + 1):
            cell = ws.cell(row=rr, column=i)
            cell.font = Font(name=BN_FONT if i in bn_cols else BASE_FONT, size=10)
            cell.fill = FILL_ME
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        for i in range(1, len(COLUMNS) + 1):
            ws.cell(row=rr, column=i).border = BOX
        ws.row_dimensions[rr].height = 28

    dv_lang = DataValidation(type="list", formula1='"bn,en,bnglish"',
                             allow_blank=True, showDropDown=False)
    dv_lang.error = "Use bn, en or bnglish"
    dv_lang.errorTitle = "Invalid language"
    ws.add_data_validation(dv_lang)
    dv_lang.add(f"B{FIRST_DATA_ROW}:B{LAST_DATA_ROW}")

    dv_type = DataValidation(type="list", formula1=f'"{",".join(QUERY_TYPES)}"',
                             allow_blank=True, showDropDown=False)
    dv_type.error = "Pick one of the listed query types"
    dv_type.errorTitle = "Invalid query_type"
    ws.add_data_validation(dv_type)
    dv_type.add(f"C{FIRST_DATA_ROW}:C{LAST_DATA_ROW}")

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{LAST_DATA_ROW}"

    # ---------------- Column guide sheet ----------------
    guide = wb.create_sheet("Column guide")
    guide["A1"] = "Column"
    guide["B1"] = "What to put in it"
    for cell in (guide["A1"], guide["B1"]):
        cell.font = HEAD_FONT
        cell.fill = HEAD_FILL
    for i, (name, helptext, _) in enumerate(COLUMNS, start=2):
        guide.cell(row=i, column=1, value=name).font = BOLD
        guide.cell(row=i, column=2, value=helptext).font = BODY
        guide.cell(row=i, column=1).border = BOX
        guide.cell(row=i, column=2).border = BOX
    r = len(COLUMNS) + 3
    guide.cell(row=r, column=1, value="query_type values").font = BOLD
    meanings = {
        "factual": "A plain fact stated in a document.",
        "date": "A deadline or schedule. Watch Bangla vs English numerals.",
        "policy": "A rule or regulation, usually a paragraph not a sentence.",
        "name": "A person, department or course name.",
        "number": "A count, credit, fee or percentage.",
        "table": "The answer lives inside a table and must come back as one.",
        "multi_hop": "Needs two or more facts combined to answer.",
        "comparison": "Needs two documents compared, e.g. 2025 vs 2026 calendar.",
        "unanswerable": "Genuinely not in any document. System must refuse.",
    }
    for i, qt in enumerate(QUERY_TYPES, start=r + 1):
        guide.cell(row=i, column=1, value=qt).font = BODY
        guide.cell(row=i, column=2, value=meanings[qt]).font = BODY
    guide.column_dimensions["A"].width = 20
    guide.column_dimensions["B"].width = 78

    wb.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
